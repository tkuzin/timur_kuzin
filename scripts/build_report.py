"""Build the report from the retained source template and measured artifacts.
Run with the bundled Codex document Python runtime (python-docx).
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
REF = Path('C:/Users/dkuzi/Downloads/Шаблон отчёта. ДЗ 1.docx')
OUT = ROOT / 'artifacts/ДЗ1_Кузин_Тимур.docx'
assert hashlib.sha256(REF.read_bytes()).hexdigest() == '230dca3b8353f412c84d6c85a6e0f7f8451ea1782b3841251b30913dc7da298d'
doc = Document(REF)
paras = list(doc.paragraphs)
small = {x['scenario']: x for x in json.loads((ROOT/'artifacts/benchmark_small.json').read_text())['results']}
work = {x['scenario']: x for x in json.loads((ROOT/'artifacts/benchmark_work.json').read_text())['results']}


def fmt(value):
    return f'{value:.2f}'.replace('.', ',')


def fill(p, value, preserve=False, code=False):
    old = deepcopy(p.runs[0]._r.rPr) if preserve and p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run = p.add_run(value)
    if old is not None:
        run._r.insert(0, old)
    elif not preserve:
        run.font.name = 'Consolas' if code else 'Times New Roman'
        run.font.size = Pt(10 if code else 12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.italic = False
        run.bold = False
        p.paragraph_format.line_spacing = 1.0 if code else 1.5
        p.paragraph_format.first_line_indent = Twips(0 if code else 709)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if code else WD_ALIGN_PARAGRAPH.JUSTIFY
        for node in list(p._p.get_or_add_pPr()):
            if node.tag in (qn('w:pBdr'), qn('w:shd')):
                node.getparent().remove(node)
    return p


def after(p, text, code=False):
    node = OxmlElement('w:p')
    if p._p.pPr is not None:
        node.append(deepcopy(p._p.pPr))
    p._p.addnext(node)
    new = Paragraph(node, p._parent)
    new.paragraph_format.page_break_before = False
    new.paragraph_format.keep_with_next = False
    return fill(new, text, code=code)


def replace_block(index, texts):
    p = fill(paras[index], texts[0])
    for text in texts[1:]:
        p = after(p, text)
    return p


def remove(index):
    node = paras[index]._p
    node.getparent().remove(node)


def table(index, headers, rows, widths):
    t = doc.tables[index]
    # Strict template: original header, grid, widths, cell and paragraph
    # formatting stay unchanged. Only fill slots and clone source body rows.
    assert len(rows[0]) == len(t.columns)
    patterns = [deepcopy(r._tr) for r in t.rows[1:]]
    if not patterns:
        pattern=deepcopy(t.rows[0]._tr)
        for node in pattern.xpath('.//w:tblHeader | .//w:b | .//w:bCs | .//w:shd'):
            node.getparent().remove(node)
        patterns=[pattern]
    while len(t.rows) < len(rows)+1:
        t._tbl.append(deepcopy(patterns[(len(t.rows)-1) % len(patterns)]))
    while len(t.rows) > len(rows)+1:
        t._tbl.remove(t.rows[-1]._tr)
    for row,values in zip(t.rows[1:],rows):
        for cell,value in zip(row.cells,values):
            p=cell.paragraphs[0]
            rpr=deepcopy(p.runs[-1]._r.rPr) if p.runs and p.runs[-1]._r.rPr is not None else None
            p.clear()
            run=p.add_run(str(value))
            if rpr is not None: run._r.insert(0,rpr)
    return t


fill(paras[13], 'Вариант 2', preserve=True)
group = os.environ.get('REPORT_GROUP', 'БИСТ-24-ПО-?')
fill(paras[14], f'Студент группы {group}\t_______________\tКузин Т. Д.', preserve=True)
remove(19)
table(0, ['Параметр','Значение'], [
    ['Вариант и предметная область','2 — заявки в службу поддержки'],
    ['Префикс','timur_kuzin'],
    ['Язык и каркас серверной части','Python 3.14.1; FastAPI 0.143.0; Uvicorn 0.54.0'],
    ['Доступ к данным','SQLAlchemy 2.1.4; psycopg 3.3.6'],
    ['СУБД','PostgreSQL 18.4'],
    ['Репозиторий',os.environ.get('REPORT_REPOSITORY','________________________________')]], [3600,6089])
fill(paras[22], 'Сервис принимает заявки пользователей с указанием категории и хранит переписку с сотрудниками. Заявка последовательно проходит статусы «новая», «в работе», «закрыта»; пропуск этапа и комментарий к закрытой заявке отклоняются. Срок реакции отсчитывается от создания до первого комментария сотрудника. Сводка показывает открытые заявки по категориям и долю нарушений норматива среди всех заявок каждой категории.')
replace_block(24, [
    'Время ответа клиент-серверного приложения зависит не только от корректности функций, но и от объёма обрабатываемых данных. На небольшом наборе полные просмотры таблиц и обработка объектов в памяти могут оставаться незаметными. Поэтому до оптимизации необходимо получить исходные показатели в воспроизводимых условиях и определить, какие операции зависят от роста базы.',
    'Цель работы — разработать сервис заявок и измерить исходные показатели его производительности на двух объёмах данных, чтобы сформулировать проверяемые гипотезы об узких местах.'
])
replace_block(25, [
    'Для достижения цели необходимо решить следующие задачи:',
    '1. Реализовать сервис с PostgreSQL, авторизацией и основными экранами.',
    '2. Описать контракт программного интерфейса.',
    '3. Покрыть бизнес-правила и операции интерфейса тестами.',
    '4. Наполнить базу малым и рабочим наборами данных воспроизводимым способом.',
    '5. Измерить время ответа всех операций на обоих объёмах.',
    '6. Разложить время замедлившихся операций на обращения к БД и оставшуюся часть.',
    '7. Сформулировать не менее трёх гипотез об узких местах.',
    '8. Проверить реализацию на очевидные ошибки и исправить их при обнаружении.'
])
replace_block(28, [
    'Интерфейс реализован шаблонами Jinja2 3.1.6 и JavaScript-запросами к API. Доступны список с фильтром и постраничным выводом, карточка с перепиской, создание заявки и сводка. Авторизация использует подписанную cookie-сессию; клиент видит свои заявки, сотрудник — все заявки и сводку.',
    'В таблице 2 Ticket означает объект с полями id, user_id, category_id, title, body, status, created_at. Формат времени — ISO 8601 с UTC. Полные схемы находятся в docs/API.md и artifacts/openapi.json. Ошибки возвращаются в поле detail. Для изменяющего запроса с чужим Origin дополнительно возвращается 403.'
])
api_rows = [
 ['POST /api/login','username, password','200: id, username, is_staff; cookie','401, 422'],
 ['POST /api/logout','Нет','200: {ok: true}','401'],
 ['GET /api/me','Нет','200: id, username, is_staff','401'],
 ['GET /api/categories','Нет','200: массив id, name, sla_minutes','401'],
 ['GET /api/tickets','page ≥ 1, size 1–100; status, category_id','200: items (Ticket[]), total, page, size','401, 422'],
 ['POST /api/tickets','title 1–200, body 1–5000, category_id > 0','201: Ticket, status=new','401, 404, 422'],
 ['GET /api/tickets/{id}','id','200: Ticket; owner, category, reaction_minutes, sla_breached, comments','401, 403, 404, 422'],
 ['PATCH /api/tickets/{id}/status','id; status','200: Ticket; new → in_progress → closed','401, 403, 404, 409, 422'],
 ['POST /api/tickets/{id}/comments','id; body 1–5000','201: id, ticket_id, body, created_at','401, 403, 404, 409, 422'],
 ['GET /api/summary','Нет; сотрудник','200: as_of, total, open, breached; categories с total, open, breached, breached_share','401, 403']
]
table(1, ['Метод и путь','Параметры','Ответ','Ошибки'], api_rows, [2850,2200,3450,1189])
fill(paras[30], 'Карточка заявки и сводка показаны на рисунках 1 и 2. Снимки сделаны после серий измерений: 50 071 заявка на сводке включает 71 запись, добавленную самим испытанием.')
p = paras[31]; p.clear(); p.paragraph_format.first_line_indent=Twips(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'artifacts/screenshots/detail.png'), width=Inches(6.45))
p.paragraph_format.keep_with_next=True
p = after(p, 'Рисунок 1 — Карточка заявки с перепиской и сроком реакции')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Twips(0)
p = after(p, ''); p.paragraph_format.page_break_before=True; p.paragraph_format.first_line_indent=Twips(0)
p.add_run().add_picture(str(ROOT/'artifacts/screenshots/summary.png'), width=Inches(6.45))
p.paragraph_format.keep_with_next=True
p = after(p, 'Рисунок 2 — Сводка по рабочему набору после измерений')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Twips(0)
replace_block(33, [
    'Команда запуска: .venv\\Scripts\\python.exe -m pytest -q. Итог: 28 passed, 1 warning in 7.50s. Предупреждение библиотеки Starlette касается будущей замены клиента httpx и не является ошибкой теста.',
    'В набор входят 11 проверок бизнес-правил без БД и 17 проверок API через TestClient с PostgreSQL. Проверены граница норматива, отсутствие ответа, первый ответ сотрудника, переходы статусов, права, фильтрация, пагинация, поля ответов и ошибки. Данные API-тестов автоматически создаются и откатываются транзакцией; запущенный веб-сервер им не требуется.'
])
code = 'def test_no_reply_has_no_reaction_time():\n    assert reaction_minutes(START, None) is None\n\ndef test_operations_require_authorization(client, method, path, payload):\n    response = client.request(method, path, json=payload)\n    assert response.status_code == 401'
listing_properties=deepcopy(paras[35]._p.pPr)
fill(paras[35], code, code=True)
paras[35]._p.remove(paras[35]._p.pPr)
paras[35]._p.insert(0,listing_properties)
for run in paras[35].runs: run.font.size=Pt(9.5)
fill(paras[36], 'Второй тест параметризован девятью защищёнными операциями. Полный набор находится в tests/. Дополнительно в браузере проверена цепочка: вход → создание → взятие в работу → ответ → закрытие; ошибок JavaScript нет.')
fill(paras[38], 'Серии выполнены 08.10.2026 на одной машине, с одной версией кода и одним процессом Uvicorn. Запросы последовательные; параллельная нагрузка не создавалась. Между объёмами изменялось только наполнение. Перед каждой серией предусмотрен прогрев.')
table(2,['Что','Значение'],[
 ['Машина','AMD Ryzen 5 5500U, 6 ядер / 12 потоков; ОЗУ 7,33 ГиБ; Windows 11, сборка 26200'],
 ['Среда выполнения стенда','Локально, без Docker; Python 3.14.1; PostgreSQL 18.4; 127.0.0.1:8080; БД localhost:55432; timur_kuzin'],
 ['Учётная запись','operator, роль сотрудника'],
 ['Число повторов в серии','30 учитываемых запросов на сценарий'],
 ['Прогрев','5 запросов перед каждой серией, исключены из расчёта'],
 ['Чем измерялось время','httpx 0.28.1 + time.perf_counter; события SQLAlchemy до и после execute; scripts/benchmark.py']],[3300,6389])
replace_block(41,[
 'Команды наполнения: .venv\\Scripts\\python.exe -m scripts.seed small и .venv\\Scripts\\python.exe -m scripts.seed work. Количества в таблице 4 получены запросами count(*) из PostgreSQL.',
 'Повторяемость проверена командой python -m scripts.verify_seed: для каждого объёма два запуска дали одинаковые количества и контрольные суммы строк каждой таблицы. Результаты сохранены в artifacts/reproducibility.json. Генератор использует фиксированные даты и формулы, после загрузки выполняет ANALYZE.'
])
table(3,['Таблица','Малое наполнение','Рабочее наполнение','Рост'],[
 ['users','50','3 000','60,00'],['categories','5','10','2,00'],['tickets','300','50 000','166,67'],['comments','900','150 000','166,67']],[1900,2600,2900,2289])
fill(paras[43], 'Чтения измерялись до операций записи. На каждом объёме испытание, включая прогрев и подготовку смены статуса, добавило 71 заявку и 35 комментариев. После малого испытания получено 371/935 записей в tickets/comments, после рабочего — 50 071/150 035. Подготовительные запросы не входят во время проверяемой операции; последующие проверки браузера выполнены после завершения замеров.')
names={'login':'Вход','logout':'Выход','me':'Профиль','categories':'Категории','list':'Список','list_filtered':'Список с фильтром','detail':'Карточка','summary':'Сводка','create':'Создание','status':'Статус','comment':'Комментарий'}
rows=[]
for key in small:
    a,b=small[key],work[key]
    rows.append([names[key],fmt(a['p50']),fmt(a['p95']),fmt(b['p50']),fmt(b['p95']),fmt(b['p50']/a['p50'])])
table(4,['Операция','p50 малое','p95 малое','p50 рабочее','p95 рабочее','Рост p50'], rows, [])
replace_block(47, [
 'Время указано в миллисекундах, рост — отношение рабочего p50 к малому. p50 — медиана, p95 — 29-е значение упорядоченной серии из 30. Список измерен без фильтра и с status=new&category_id=1; всего 10 операций API и 11 сценариев. Сырые 660 наблюдений сохранены в artifacts/samples_small.csv и samples_work.csv.',
 'Максимальное время ответа (малое / рабочее, мс): '+ '; '.join(f'{names[k].lower()} — {fmt(small[k]["max"])} / {fmt(work[k]["max"])}' for k in small)+'.'
])
replace_block(48,[
 'Наибольший рост у сводки: 19,31 → 1 352,19 мс, в 70,04 раза. Список вырос в 3,44 раза, список с фильтром — в 3,57, карточка — в 2,78. При росте основной таблицы в 166,67 раза ни одна операция не показала сверхлинейного роста между этими двумя точками. Для дальнейшего анализа выбраны перечисленные четыре сценария как зависимые от объёма и наиболее заметно замедлившиеся.',
 'Вход, выход, профиль, категории, создание, смена статуса и добавление комментария исключены из дальнейшего разбора: их медианы существенно не увеличились. У комментария вырос p95, поэтому отдельные хвостовые задержки следует проверить в следующем нагрузочном испытании; 30 повторов недостаточно для вывода о стабильности хвоста.'
])
fill(paras[50], 'SQL-время измерено событиями SQLAlchemy вокруг вызовов драйвера, число запросов — теми же событиями. Таблица 6 содержит медианы рабочего объёма; остаток рассчитан как p50 HTTP минус p50 SQL. Это оценка разложения, а не отдельный замер процессорного времени Python.')
table(5,['Операция','Всего, мс','SQL, мс','Остаток, мс','Запросов','Вывод'],[
 [names[k],fmt(work[k]['p50']),fmt(work[k]['sql_p50']),fmt(work[k]['residual']),str(int(work[k]['queries'])),'Приложение' if k=='summary' else 'БД'] for k in ('list','list_filtered','detail','summary')],[1900,1450,1450,1650,1350,1889])
fill(paras[52], 'SQL-время включает выполнение прикладных запросов, локальный обмен и часть работы драйвера. Счётчик учитывает прикладные SQL-команды; служебная проверка соединения пулом и управление транзакциями отдельно не учитываются. Остаток включает эти накладные расходы, построение ORM-объектов, Python, сериализацию и HTTP. Чистое время внутри PostgreSQL и CPU приложения на этом этапе отдельно не профилировались.')
fill(paras[53], 'У списка около 76% медианы приходится на обращения к БД, у списка с фильтром — 77%, у карточки — 72%. У сводки преобладает остаток: 1 094,80 мс, около 81% общего времени. Число запросов не растёт с объёмом: 3 у списка, 5 у карточки и 4 у сводки; признаков N+1 в этих операциях не обнаружено.')
remove(55)
replace_block(56,[
 'Гипотеза 1. Сводка и обработка всех заявок в приложении.',
 '1. Операция: GET /api/summary.',
 '2. Наблюдение: p50 вырос с 19,31 до 1 352,19 мс; SQL — 257,39 мс, остаток — 1 094,80 мс, запросов — 4.',
 '3. Предполагаемая причина: приложение получает все 50 000 заявок, создаёт ORM-объекты и считает показатели по категориям циклом Python. Проверка: профиль затрат на материализацию объектов и агрегирование; затем сравнение с агрегатами SQL при неизменном ответе API.'
])
replace_block(57,[
 'Гипотеза 2. Подсчёт и сортировка списка без дополнительных индексов.',
 '1. Операция: GET /api/tickets, без фильтра и с фильтром.',
 '2. Наблюдение: p50 42,93 и 46,98 мс; SQL — 32,49 и 35,99 мс, по 3 запроса. На малом наборе p50 составлял 12,48 и 13,16 мс.',
 '3. Предполагаемая причина: count(*) и выборка с фильтрами/сортировкой просматривают значительную часть таблицы. Дополнительных индексов нет. Проверка: EXPLAIN (ANALYZE, BUFFERS) для подсчёта и страницы; индекс и изменение запроса допускаются только на этапе оптимизации.'
])
replace_block(58,[
 'Гипотеза 3. Поиск комментариев заявки в большой таблице.',
 '1. Операция: GET /api/tickets/{id}.',
 '2. Наблюдение: p50 вырос с 13,50 до 37,53 мс, SQL — с 4,07 до 26,98 мс. В обоих случаях 5 запросов, у выбранной заявки 3 комментария.',
 '3. Предполагаемая причина: выборка комментариев по ticket_id просматривает таблицу из 150 000 строк. Внешний ключ PostgreSQL сам по себе не создаёт индекс на ссылающемся столбце; в исходной схеме есть только первичные индексы. Проверка: план и число просмотренных строк запроса комментариев.'
])
fill(paras[59], 'Очевидных ошибок реализации не обнаружено: 28 тестов проходят, основной сценарий проверен в браузере, пагинация ограничивает выдачу. По коду и счётчикам SQL повторных одинаковых запросов нет. Ошибки Windows-скрипта запуска исправлены до обеих серий. Оптимизация пока не выполнялась.')
for idx in (56,57,58):
    for r in paras[idx].runs: r.bold=True
replace_block(61,[
 'Цель работы достигнута: сервис заявок реализован, исходные показатели получены на двух объёмах и позволяют выбрать направления дальнейшего исследования.',
 'Для этого решены следующие задачи:',
 '1. Реализован сервис с 4 сущностями, PostgreSQL, авторизацией и тремя обязательными экранами.',
 '2. Описан контракт 10 операций программного интерфейса.',
 '3. Выполнены 28 проверок бизнес-правил и API, все прошли.',
 '4. Подготовлены воспроизводимые наборы с 300 и 50 000 заявок; рабочий набор содержит 150 000 комментариев.',
 '5. Измерены 11 сценариев на двух объёмах, сохранены 660 учитываемых наблюдений.',
 '6. Для четырёх замедлившихся сценариев оценены SQL-время, остаток и число запросов.',
 '7. Сформулированы три гипотезы об обработке данных сводки, запросах списка и поиске комментариев.',
 '8. Проверено отсутствие очевидных функциональных ошибок; исправлены проблемы локального запуска до измерений.'
])
fill(paras[62], 'Основное узкое место — сводка: p50 составляет 1 352,19 мс, из которых 1 094,80 мс приходится на остаток вне измеренных SQL-вызовов. У списка и карточки преобладают обращения к БД. Эти выводы определяют направления профилирования; причинность и эффект изменений предстоит подтвердить на следующих этапах, сохранив контракт API и условия испытания.')

# Preserve package parts unrelated to authored body, images and field settings.
for node in doc.element.xpath('//w:t'):
    for number in ('1 352','1 094','50 000','150 000','50 071','150 035'):
        if node.text: node.text=node.text.replace(number,number.replace(' ','\u00a0'))
update = OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); doc.settings.element.append(update)
buffer=BytesIO(); doc.save(buffer)
editable={'word/document.xml','word/_rels/document.xml.rels','word/settings.xml','[Content_Types].xml'}
with ZipFile(REF) as source, ZipFile(BytesIO(buffer.getvalue())) as built, ZipFile(OUT,'w',ZIP_DEFLATED) as out:
    for info in source.infolist():
        out.writestr(info,built.read(info.filename) if info.filename in editable else source.read(info.filename))
    for name in built.namelist():
        if name not in source.namelist(): out.writestr(name,built.read(name))
with ZipFile(REF) as source, ZipFile(OUT) as final:
    assert all(source.read(n)==final.read(n) for n in source.namelist() if n not in editable)
assert hashlib.sha256(REF.read_bytes()).hexdigest() == '230dca3b8353f412c84d6c85a6e0f7f8451ea1782b3841251b30913dc7da298d'
print(OUT)
